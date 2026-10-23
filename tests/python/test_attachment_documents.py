"""Synthetic fixtures only. Real conversion: TM_DOCLING_TEST_REAL=1, TM_DOCLING_MODELS=<models>."""
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

spec = importlib.util.spec_from_file_location("bridge", Path(__file__).parents[2] / "src/main/attachments/docling-bridge.py")
bridge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bridge)


def pdf(path):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>"]
    for content_id, label in ((4, b"Budget approval: 12345 dollars."), (6, b"Delivery deadline: November 30.")):
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents {content_id} 0 R >>".encode())
        stream = b"BT /F1 16 Tf 72 700 Td (" + label + b") Tj ET"
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    data = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(data)
    data += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode()
    data += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    path.write_bytes(data)


class AttachmentIndexTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = dict(conversationId="group-a", messageId="local:7", createTime=100, fileName="synthetic.pdf", sender="测试成员", senderId="sender-a")
        self.file = self.root / "synthetic.pdf"
        pdf(self.file)

    def tearDown(self):
        self.temp.cleanup()

    def call(self, op, **kwargs):
        return bridge.main(dict(root=str(self.root), accountId="test-account", models=os.environ.get("TM_DOCLING_MODELS", ""), op=op, **kwargs))

    def index(self, source=None, text="附件预算报价 12345，交付期限 November"):
        with patch.object(bridge, "convert", return_value=[dict(location="第 1 页", text=text)]):
            return self.call("parse", source=source or self.source, path=str(self.file))

    def test_search_scope_provenance_and_short_chinese(self):
        self.index()
        for term in ("预算", "12345", "November"):
            rows = self.call("search", terms=[term], conversationIds=["group-a"], startTime=100, endTime=100)["evidence"]
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["derivedSource"], "attachment_text")
            self.assertEqual(rows[0]["timestamp"], 100000)
            self.assertEqual(rows[0]["messageId"], "local:7")
            self.assertIn("第 1 页", rows[0]["text"])
        for scope in (dict(conversationIds=["other"]), dict(senderIds=["other"]), dict(startTime=101), dict(endTime=99)):
            self.assertEqual(self.call("search", terms=["预算"], **scope)["evidence"], [])
        self.assertEqual(self.call("search", terms=["%_"])["evidence"], [])
        self.assertEqual(len(self.call("search", terms=["预算"], conversationIds=[f"noise-{i}" for i in range(800)] + ["group-a"])["evidence"]), 1)

    def test_account_isolation_read_stale_and_delete(self):
        self.index()
        self.assertEqual(bridge.main(dict(root=str(self.root), accountId="other", op="search", terms=["预算"]))["evidence"], [])
        self.assertEqual(len(self.call("read", conversationId="group-a", messageIds=["local:7"])["documents"]), 1)
        self.file.write_bytes(self.file.read_bytes() + b"\nchanged")
        self.assertEqual(self.call("status", source=self.source)["state"], "stale")
        self.assertEqual(self.call("search", terms=["预算"])["evidence"], [])
        self.assertEqual(self.call("read", conversationId="group-a", messageIds=["local:7"])["documents"], [])
        self.call("delete", source=self.source)
        self.assertEqual(self.call("status", source=self.source)["state"], "unparsed")
        self.assertTrue(self.file.exists())

    def test_long_attachment_does_not_crowd_out_other_source_messages(self):
        self.index(text="预算" * 20)
        self.index(source={**self.source, "messageId": "local:8"}, text="预算来自另一份文件")
        db = bridge.connect(dict(root=str(self.root), accountId="test-account"))
        with db:
            for number in range(10):
                db.execute("INSERT INTO parts VALUES (?,?,?)", (bridge.document_key(self.source), f"第 {number+2} 页", "预算" * 30))
        db.close()
        hits = self.call("search", terms=["预算"], limit=2)["evidence"]
        self.assertEqual({hit["messageId"] for hit in hits}, {"local:7", "local:8"})
        self.assertLessEqual(sum(hit["messageId"] == "local:7" for hit in hits), 3)

    def test_scoping_happens_before_candidate_limit(self):
        db = bridge.connect(dict(root=str(self.root), accountId="test-account"))
        with db:
            for i in range(1002):
                source = {**self.source, "conversationId": "target" if i == 1001 else "noise", "messageId": str(i)}
                key = bridge.document_key(source)
                db.execute("INSERT INTO documents VALUES (?,?,?,?,?,?,?,?,?)", (key, source["conversationId"], str(i), json.dumps(source), str(self.file), bridge.fingerprint(self.file), "test", 1, bridge.VERSION))
                db.execute("INSERT INTO parts VALUES (?,?,?)", (key, "第 1 页", "uniqueBudget"))
        db.close()
        self.assertEqual(len(self.call("search", terms=["uniqueBudget"], conversationIds=["target"])["evidence"]), 1)

    def test_invalid_and_failed_parse_never_write_partial_document(self):
        self.file.write_bytes(b"not a pdf")
        with self.assertRaisesRegex(ValueError, "PDF"):
            self.index()
        self.assertEqual(self.call("status", source=self.source)["state"], "unparsed")
        pdf(self.file)
        with patch.object(bridge, "convert", side_effect=ValueError("conversion failed")):
            with self.assertRaisesRegex(ValueError, "conversion failed"):
                self.call("parse", source=self.source, path=str(self.file))
        self.assertEqual(self.call("status", source=self.source)["state"], "unparsed")

    @unittest.skipUnless(os.environ.get("TM_DOCLING_TEST_REAL") == "1", "Opt-in local Docling conversion")
    def test_real_docling_pdf_docx_xlsx(self):
        from docx import Document
        from openpyxl import Workbook
        word = self.root / "synthetic.docx"
        doc = Document(); doc.add_heading("项目计划", 0); doc.add_paragraph("附件正文中文验证：预算为 67890 元。")
        table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "交付"; table.cell(0, 1).text = "十一月"
        doc.save(word)
        excel = self.root / "synthetic.xlsx"
        book = Workbook(); book.active.title = "预算"; book.active.append(["项目", "金额"]); book.active.append(["研发", 24680])
        sheet = book.create_sheet("排期"); sheet.append(["阶段", "日期"]); sheet.append(["验收", "2026-11-30"]); book.save(excel)
        for number, (file, term, location) in enumerate(((self.file, "12345", "第 1 页"), (word, "67890", "段落/表格"), (excel, "24680", "预算"))):
            source = {**self.source, "messageId": str(number), "fileName": file.name}
            result = self.call("parse", source=source, path=str(file))
            self.assertEqual(result["state"], "indexed")
            hits = self.call("search", terms=[term])["evidence"]
            self.assertTrue(any(term in hit["text"] and location in hit["text"] for hit in hits), hits)
        hits = self.call("search", terms=["November"])["evidence"]
        self.assertTrue(any("第 2 页" in hit["text"] for hit in hits))
        hits = self.call("search", terms=["2026-11-30"])["evidence"]
        self.assertTrue(any("排期" in hit["text"] for hit in hits))

    @unittest.skipUnless(os.environ.get("TM_DOCLING_TEST_REAL") == "1", "Opt-in local Docling conversion")
    def test_scanned_pdf_ocr_runs_with_local_models(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1200, 500), "white")
        font = ImageFont.load_default(size=48)
        ImageDraw.Draw(image).text((70, 160), "Budget approval 54321 dollars", fill="black", font=font)
        image.save(self.file, "PDF", resolution=144)
        self.call("parse", source=self.source, path=str(self.file))
        hits = self.call("search", terms=["54321"])["evidence"]
        self.assertTrue(any("第 1 页" in hit["text"] for hit in hits), hits)


if __name__ == "__main__":
    unittest.main()
