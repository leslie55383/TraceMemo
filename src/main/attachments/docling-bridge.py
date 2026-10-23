"""Local document conversion and account-scoped FTS. JSON stdin/stdout; no chat database access."""
import contextlib
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import sqlite3
import sys
import time
import zipfile

VERSION = "2.142.0"
OCR_VERSION = "1.31.0"
MAX_CHARACTERS = 1_000_000
MAX_PARTS = 1500


def fingerprint(path):
    stat = Path(path).stat()
    return f"{stat.st_size}:{stat.st_mtime_ns}"


def connect(request):
    account = request.get("accountId", "")
    if not isinstance(account, str) or not account:
        raise ValueError("没有已连接的微信账号")
    folder = Path(request["root"]) / "attachment-text" / hashlib.sha256(account.encode()).hexdigest()
    folder.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(folder / "documents.sqlite", timeout=15)
    db.row_factory = sqlite3.Row
    db.executescript("""
      PRAGMA journal_mode=WAL;
      CREATE TABLE IF NOT EXISTS documents (
        id TEXT PRIMARY KEY, conversation_id TEXT NOT NULL, message_id TEXT NOT NULL,
        source_json TEXT NOT NULL, path TEXT NOT NULL, signature TEXT NOT NULL,
        content_hash TEXT NOT NULL, parsed_at INTEGER NOT NULL, engine TEXT NOT NULL,
        UNIQUE(conversation_id, message_id));
      CREATE VIRTUAL TABLE IF NOT EXISTS parts USING fts5(
        document_id UNINDEXED, location UNINDEXED, text, tokenize='trigram');
    """)
    return db


def document_key(source):
    return hashlib.sha256(json.dumps([source["conversationId"], source["messageId"]], ensure_ascii=False).encode()).hexdigest()


def fresh(row):
    try:
        return row["signature"] == fingerprint(row["path"]) and row["engine"] == VERSION
    except OSError:
        return False


def document(db, row):
    return {
        "source": json.loads(row["source_json"]), "parsedAt": row["parsed_at"], "engine": row["engine"],
        "parts": [dict(p) for p in db.execute("SELECT location, text FROM parts WHERE document_id=? ORDER BY rowid", (row["id"],))]
    }


def validate_file(path):
    path = Path(path).resolve(strict=True)
    if not path.is_file() or path.suffix.lower() not in (".pdf", ".docx", ".xlsx"):
        raise ValueError("仅支持 PDF、DOCX、XLSX 本地文件")
    if path.stat().st_size > 30 * 1024 * 1024:
        raise ValueError("附件超过 30 MB，请拆分后处理")
    if path.suffix.lower() == ".pdf":
        with path.open("rb") as stream:
            if not stream.read(1024).lstrip().startswith(b"%PDF-"):
                raise ValueError("文件内容不是有效 PDF")
    else:
        with zipfile.ZipFile(path) as archive:
            if sum(p.file_size for p in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError("Office 文件解压后超过 150 MB，请拆分处理")
            expected = "word/document.xml" if path.suffix.lower() == ".docx" else "xl/workbook.xml"
            if expected not in archive.namelist():
                raise ValueError("Office 文件内容与扩展名不匹配")
    return path


def convert(path, models):
    from docling.datamodel.base_models import InputFormat, ConversionStatus
    from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    options = PdfPipelineOptions(artifacts_path=Path(models), enable_remote_services=False, allow_external_plugins=False)
    options.ocr_options = RapidOcrOptions(lang=["ch"])
    converter = DocumentConverter(allowed_formats=[InputFormat.PDF, InputFormat.DOCX, InputFormat.XLSX],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})
    # No network while handling a document, including embedded remote resources/model fallback.
    import socket
    def deny_network(*args, **kwargs):
        raise RuntimeError("文档解析禁止联网；请先准备本地模型")
    socket.socket.connect = deny_network
    socket.socket.connect_ex = deny_network
    result = converter.convert(path, max_num_pages=500, max_file_size=30 * 1024 * 1024)
    if result.status != ConversionStatus.SUCCESS:
        raise ValueError("附件未完整转换，未写入索引；请检查加密、损坏或模型状态")
    doc = result.document
    sheet_names = []
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheet_names = list(book.sheetnames)
        book.close()
    parts = []
    total = 0
    for number, (item, _) in enumerate(doc.iterate_items(), 1):
        text = getattr(item, "text", "")
        if hasattr(item, "data") and hasattr(item, "export_to_markdown"):
            text = item.export_to_markdown(doc=doc)
        if not text or not text.strip():
            continue
        pages = sorted({p.page_no for p in getattr(item, "prov", [])})
        if sheet_names and pages:
            location = "工作表：" + "、".join(sheet_names[p - 1] for p in pages if 0 < p <= len(sheet_names))
        elif path.suffix.lower() == ".pdf" and pages:
            location = "第 " + "、".join(map(str, pages)) + " 页"
        else:
            location = f"段落/表格 {number}"  # DOCX does not have stable rendered page numbers.
        for offset in range(0, len(text), 2000):
            piece = text[offset:offset + 2000]
            total += len(piece)
            parts.append({"location": location, "text": piece})
            if total > MAX_CHARACTERS or len(parts) > MAX_PARTS:
                raise ValueError("正文超过单份附件索引上限，未截断入库；请拆分文件")
    if not parts:
        raise ValueError("附件没有可提取的正文；扫描件请确认 OCR 模型已准备")
    return parts


def main(request):
    op = request["op"]
    if op == "probe":
        try:
            version = importlib.metadata.version("docling")
            ocr_ready = importlib.metadata.version("onnxruntime") == OCR_VERSION
        except importlib.metadata.PackageNotFoundError:
            version = ""
            ocr_ready = False
        marker = Path(request["models"]) / "tracememo-ready.json"
        try:
            pdf_ready = json.loads(marker.read_text()).get("version") == VERSION
        except (OSError, ValueError):
            pdf_ready = False
        return {"ready": version == VERSION and ocr_ready, "version": version, "pdfReady": pdf_ready and ocr_ready, "preparing": False}
    if op == "prepare":
        from docling.utils.model_downloader import download_models
        download_models(Path(request["models"]), with_code_formula=False, with_picture_classifier=False, with_easyocr=False)
        (Path(request["models"]) / "tracememo-ready.json").write_text(json.dumps({"version": VERSION}))
        return {"ready": True, "version": VERSION, "pdfReady": True, "preparing": False}
    db = connect(request)
    try:
        if op in ("status", "delete", "parse"):
            source = request["source"]
            key = document_key(source)
            row = db.execute("SELECT * FROM documents WHERE id=?", (key,)).fetchone()
            if op == "delete":
                with db:
                    db.execute("DELETE FROM parts WHERE document_id=?", (key,))
                    db.execute("DELETE FROM documents WHERE id=?", (key,))
                return {"state": "unparsed"}
            if op == "status":
                if not row:
                    return {"state": "unparsed"}
                values = db.execute("SELECT count(*), coalesce(sum(length(text)),0) FROM parts WHERE document_id=?", (key,)).fetchone()
                return {"state": "indexed" if fresh(row) else "stale", "fileName": json.loads(row["source_json"])["fileName"], "partCount": values[0], "characters": values[1]}
            path = validate_file(request["path"])
            signature = fingerprint(path)
            if row and fresh(row) and row["path"] == str(path):
                return {"state": "indexed", "fileName": source["fileName"], "partCount": len(document(db, row)["parts"])}
            parts = convert(path, request["models"])
            if signature != fingerprint(path):
                raise ValueError("解析期间文件发生变化，未保存，请重试")
            with db:
                db.execute("DELETE FROM parts WHERE document_id=?", (key,))
                db.execute("INSERT OR REPLACE INTO documents VALUES (?,?,?,?,?,?,?,?,?)", (key, source["conversationId"], source["messageId"], json.dumps(source, ensure_ascii=False), str(path), signature, hashlib.sha256(path.read_bytes()).hexdigest(), int(time.time()*1000), VERSION))
                db.executemany("INSERT INTO parts VALUES (?,?,?)", [(key, part["location"], part["text"]) for part in parts])
            return {"state": "indexed", "fileName": source["fileName"], "partCount": len(parts), "characters": sum(len(p["text"]) for p in parts)}
        if op == "read":
            ids = request.get("messageIds", [])[:300]
            if not ids:
                return {"documents": []}
            rows = db.execute("SELECT * FROM documents WHERE conversation_id=? AND message_id IN (" + ",".join("?" * len(ids)) + ")", [request["conversationId"], *ids]).fetchall()
            docs = [document(db, row) for row in rows if fresh(row)]
            if sum(len(part["text"]) for doc in docs for part in doc["parts"]) > MAX_CHARACTERS:
                raise ValueError("所选附件正文合计超过 100 万字符，请缩小日报范围")
            return {"documents": docs}
        if op == "search":
            terms = [str(t).strip()[:100] for t in request.get("terms", [])[:12] if str(t).strip()]
            if not terms:
                terms = [str(request.get("text", "")).strip()[:100]]
            terms = [t for t in terms if t]
            if not terms:
                return {"evidence": []}
            clauses, args = [], []
            for term in terms:
                if len(term) >= 3:
                    clauses.append("parts MATCH ?")
                    args.append('"' + term.replace('"', '""') + '"')
                else:
                    clauses.append("parts.text LIKE ? ESCAPE '\\'")
                    args.append("%" + term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%")
            # Separate MATCH queries avoid SQLite's MATCH-in-OR limitation.
            filters, filter_args = [], []
            for field, key in (("conversation_id", "conversationIds"), ("json_extract(source_json,'$.senderId')", "senderIds")):
                values = request.get(key, [])
                if values:
                    filters.append(field + " IN (SELECT value FROM json_each(?))")
                    filter_args.append(json.dumps(values))
            for key, comparison in (("startTime", ">="), ("endTime", "<=")):
                if request.get(key) is not None:
                    filters.append("json_extract(source_json,'$.createTime') " + comparison + " ?")
                    filter_args.append(request[key])
            scope = "".join(" AND " + f for f in filters)
            matched = {}
            for clause, arg in zip(clauses, args):
                rows = db.execute("SELECT documents.*, parts.rowid AS part_id, parts.location, parts.text FROM parts JOIN documents ON documents.id=parts.document_id WHERE " + clause + scope + " ORDER BY json_extract(source_json,'$.createTime') DESC LIMIT 1000", [arg, *filter_args]).fetchall()
                for row in rows:
                    matched[row["part_id"]] = row
            evidence = []
            for row in matched.values():
                source = json.loads(row["source_json"])
                stamp = source["createTime"] * 1000
                if request.get("conversationIds") and source["conversationId"] not in request["conversationIds"]:
                    continue
                if request.get("senderIds") and source.get("senderId") not in request["senderIds"]:
                    continue
                if request.get("startTime") is not None and stamp < request["startTime"] * 1000:
                    continue
                if request.get("endTime") is not None and stamp > request["endTime"] * 1000:
                    continue
                if not fresh(row):
                    continue
                text = f"附件《{source['fileName']}》 · {row['location']}（本地解析正文，非聊天原文）\n{row['text']}"
                evidence.append({"chunkId": "attachment:" + str(row["part_id"]), "conversationId": source["conversationId"], "messageId": source["messageId"], "messageIds": [source["messageId"]], "senderId": source.get("senderId"), "sender": source["sender"], "timestamp": stamp, "startTime": stamp, "endTime": stamp, "sourceKind": "file", "derivedSource": "attachment_text", "text": text, "score": sum(row["text"].lower().count(t.lower()) for t in terms) + 1})
            evidence.sort(key=lambda x: (-x["score"], -x["timestamp"]))
            # `limit` counts source messages, not chunks: a long PDF must not crowd out other files.
            limit = max(1, min(int(request.get("limit", 20)), 100))
            selected, per_document = [], {}
            for item in evidence:
                key = (item["conversationId"], item["messageId"])
                if key not in per_document and len(per_document) >= limit:
                    continue
                count = per_document.get(key, 0)
                if count < 3:
                    selected.append(item)
                    per_document[key] = count + 1
            return {"evidence": selected}
        raise ValueError("不支持的附件操作")
    finally:
        db.close()


if __name__ == "__main__":
    try:
        request = json.load(sys.stdin)
        # Third-party parsers must not contaminate the JSON protocol with progress output.
        with contextlib.redirect_stdout(sys.stderr):
            result = main(request)
        print(json.dumps({"success": True, "result": result}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"success": False, "error": str(error)[:600]}, ensure_ascii=False))
